"""
Build the MoodSyncAI architecture diagram.

Produces TWO outputs from a single shared specification:

    assets/architecture.png             — high-resolution PNG (matplotlib)
    assets/architecture_diagram.pptx    — fully editable PowerPoint version

Both follow the same clean, light, four-layer layered style:
    Layer 1  PERCEPTION  — input capture and pre-processing      (blue)
    Layer 2  COGNITION   — modality models, fusion, mismatch     (green)
    Layer 3  SYNTHESIS   — FLAN-T5 summary + Conversation Coach  (purple)
    Layer 4  INTERFACE   — Streamlit UI + session state          (yellow)

Run from project root:
    python scripts/build_architecture.py
"""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt


# ----------------------------------------------------------------------
# Shared spec — single source of truth for both PNG and PPTX renders
# ----------------------------------------------------------------------

TITLE = "MoodSyncAI — Four-Layer Inference Pipeline"

# Layer palette (light fill + matching border) — clean, business-like.
LAYERS = {
    "perception": {
        "label":   "LAYER 1  ·  PERCEPTION",
        "fill":    "#EFF5FE",   # pale blue
        "border":  "#3B82F6",   # blue
        "title":   "Input capture & pre-processing",
    },
    "cognition": {
        "label":   "LAYER 2  ·  COGNITION",
        "fill":    "#EDFBF3",   # pale green
        "border":  "#10B981",   # emerald
        "title":   "Modality models + fusion",
    },
    "synthesis": {
        "label":   "LAYER 3  ·  SYNTHESIS",
        "fill":    "#F4EFFE",   # pale violet
        "border":  "#8B5CF6",   # violet
        "title":   "Summary + Coach advice",
    },
    "interface": {
        "label":   "LAYER 4  ·  INTERFACE",
        "fill":    "#FEF6E0",   # pale amber
        "border":  "#F59E0B",   # amber
        "title":   "Streamlit UI + state",
    },
}

# Per-layer "PROCESSING" inner-card contents.
# For PowerPoint we render with emoji icons (good visual weight in the deck).
# For the matplotlib PNG we substitute a plain bullet so the embedded
# diagram renders cleanly without depending on an emoji-capable font.
PERCEPTION_NODES = [
    ("🖼", "Image / Photo"),
    ("📸", "Webcam capture"),
    ("✍", "Text input"),
    ("🎙", "Audio clip"),
    ("🎥", "Video upload"),
]
PERCEPTION_PIPELINE = [
    "PIL → RGB",
    "AutoImageProcessor 224×224",
    "RoBERTa BPE tokeniser (≤128 tok)",
    "Whisper feature extractor",
]

COGNITION_PIPELINE = [
    ("🧠", "ViT face emotion"),
    ("📝", "DistilRoBERTa text"),
    ("🔊", "Whisper ASR"),
    ("⚖", "Polarity reduction"),
    ("🔀", "Fusion MLP / Weighted"),
    ("⚠", "Mismatch L2 score"),
]

SYNTHESIS_PIPELINE = [
    ("✨", "FLAN-T5 summary"),
    ("📈", "Trajectory detection (7 patterns)"),
    ("🧭", "Conversation Coach advice"),
    ("🧩", "Template fallback"),
]

INTERFACE_PIPELINE = [
    ("📷", "Image + Text tab"),
    ("🎥", "Webcam / Video tab"),
    ("🎙", "Audio + Image tab"),
    ("🧭", "Conversation Coach tab"),
    ("🎨", "Glass-card UI + Plotly charts"),
    ("💾", "session_state across turns"),
]


def _strip_icon(item):
    """Drop the emoji icon for matplotlib rendering — leaves a clean bullet."""
    if isinstance(item, tuple):
        return ("•", item[1])
    return item


# ======================================================================
# 1)  matplotlib PNG renderer
# ======================================================================

ROOT     = Path(__file__).resolve().parent.parent
OUT_PNG  = ROOT / "assets" / "architecture.png"
OUT_PNG2 = ROOT / "docs"   / "architecture.png"
OUT_PPTX = ROOT / "assets" / "architecture_diagram.pptx"


def _round_box(ax, x, y, w, h, *, fc, ec, lw=1.8, alpha=1.0, zorder=2,
               radius=0.18):
    box = FancyBboxPatch(
        (x, y), w, h,
        boxstyle=f"round,pad=0.0,rounding_size={radius}",
        linewidth=lw, edgecolor=ec, facecolor=fc, alpha=alpha,
        zorder=zorder,
    )
    ax.add_patch(box)
    return box


def _add_layer_card(ax, x, y, w, h, *, spec, zorder=2):
    # Outer panel (fill + border)
    _round_box(ax, x, y, w, h,
               fc=spec["fill"], ec=spec["border"], lw=2.0,
               zorder=zorder)
    # Label tag (small pill in the top-left)
    ax.text(
        x + 0.30, y + h - 0.40,
        spec["label"],
        fontsize=9, weight="bold", color=spec["border"],
        family="DejaVu Sans", zorder=zorder + 1,
    )
    # Section title
    ax.text(
        x + 0.30, y + h - 0.85,
        spec["title"],
        fontsize=12.5, weight="bold", color="#1A1F2C",
        family="DejaVu Sans", zorder=zorder + 1,
    )


def _add_processing_card(ax, x, y, w, h, *, items, zorder=3):
    """White inner card with a vertical list of (icon, label) rows."""
    _round_box(ax, x, y, w, h,
               fc="#FFFFFF", ec="#C8CFDA", lw=1.0,
               zorder=zorder, radius=0.14)
    ax.text(
        x + 0.18, y + h - 0.30,
        "PROCESSING",
        fontsize=8, weight="bold", color="#8A94A6",
        family="DejaVu Sans", zorder=zorder + 1,
    )
    # Stack of rows
    n = len(items)
    row_h = (h - 0.6) / max(n, 1)
    for i, item in enumerate(items):
        row_y = y + h - 0.65 - (i + 1) * row_h + 0.05
        if isinstance(item, tuple):
            icon, label = item
            ax.text(x + 0.30, row_y + row_h * 0.45, icon,
                    fontsize=11, color="#374151",
                    family="DejaVu Sans", zorder=zorder + 1)
            ax.text(x + 0.85, row_y + row_h * 0.45, label,
                    fontsize=10, color="#1A1F2C",
                    family="DejaVu Sans", zorder=zorder + 1)
        else:
            ax.text(x + 0.30, row_y + row_h * 0.45, item,
                    fontsize=10, color="#1A1F2C",
                    family="DejaVu Sans", zorder=zorder + 1)


def _add_input_node(ax, x, y, w, h, icon, label, *, accent="#3B82F6"):
    _round_box(ax, x, y, w, h, fc="#FFFFFF", ec=accent, lw=1.4,
               zorder=4, radius=0.10)
    ax.text(x + 0.18, y + h * 0.5, icon,
            fontsize=12, color=accent, family="DejaVu Sans",
            verticalalignment="center", zorder=5)
    ax.text(x + 0.55, y + h * 0.5, label,
            fontsize=10, color="#1A1F2C", family="DejaVu Sans",
            verticalalignment="center", zorder=5)


def _arrow(ax, p1, p2, *, color="#374151", lw=1.6, style="-|>"):
    ax.add_patch(FancyArrowPatch(
        p1, p2,
        arrowstyle=style, mutation_scale=18,
        linewidth=lw, color=color, zorder=6,
        shrinkA=4, shrinkB=4,
    ))


def render_png():
    """High-res PNG via matplotlib."""
    fig_w, fig_h = 20, 10
    fig = plt.figure(figsize=(fig_w, fig_h), dpi=180, facecolor="#FFFFFF")
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, fig_w)
    ax.set_ylim(0, fig_h)
    ax.axis("off")

    # Title
    ax.text(fig_w / 2, 9.55, TITLE,
            fontsize=20, weight="bold", color="#1A1F2C",
            ha="center", family="DejaVu Sans")
    ax.text(fig_w / 2, 9.15,
            "Image + Text + Audio + Video  →  fusion + mismatch  →  summary + coach  →  UI",
            fontsize=11, color="#4A5568",
            ha="center", family="DejaVu Sans", style="italic")

    # ===== Input nodes (left edge) =====
    inputs = [_strip_icon(item) for item in PERCEPTION_NODES]
    for i, (icon, label) in enumerate(inputs):
        y = 6.8 - i * 0.85
        _add_input_node(ax, 0.4, y, 2.4, 0.65, icon, label,
                        accent=LAYERS["perception"]["border"])

    # ===== LAYER 1  PERCEPTION (blue, left) =====
    L1 = LAYERS["perception"]
    L1_BOX = (3.2, 1.2, 3.6, 7.4)  # x, y, w, h
    _add_layer_card(ax, *L1_BOX, spec=L1)
    _add_processing_card(ax, 3.45, 1.6, 3.1, 4.4,
                         items=[_strip_icon(it) for it in PERCEPTION_PIPELINE])

    # Arrows from inputs into layer 1
    for i in range(len(inputs)):
        y = 7.13 - i * 0.85
        _arrow(ax, (2.8, y), (3.4, y), color=L1["border"], lw=1.2)

    # ===== LAYER 2  COGNITION (green, middle) =====
    L2 = LAYERS["cognition"]
    L2_BOX = (7.4, 1.2, 3.6, 7.4)
    _add_layer_card(ax, *L2_BOX, spec=L2)
    _add_processing_card(ax, 7.65, 1.6, 3.1, 5.4,
                         items=[_strip_icon(it) for it in COGNITION_PIPELINE])

    # Big arrow Layer 1 → Layer 2
    _arrow(ax, (6.85, 4.0), (7.4, 4.0), color="#374151", lw=2.4)

    # ===== LAYER 3  SYNTHESIS (purple, right-top) =====
    L3 = LAYERS["synthesis"]
    L3_BOX = (11.6, 5.2, 3.6, 3.4)
    _add_layer_card(ax, *L3_BOX, spec=L3)
    _add_processing_card(ax, 11.85, 5.45, 3.1, 2.5,
                         items=[_strip_icon(it) for it in SYNTHESIS_PIPELINE])

    # ===== LAYER 4  INTERFACE (yellow, right-bottom) =====
    L4 = LAYERS["interface"]
    L4_BOX = (11.6, 1.2, 3.6, 3.6)
    _add_layer_card(ax, *L4_BOX, spec=L4)
    _add_processing_card(ax, 11.85, 1.4, 3.1, 3.0,
                         items=[_strip_icon(it) for it in INTERFACE_PIPELINE])

    # ===== Cognition → result-event hub → Synthesis & Interface =====
    # Hub sits BETWEEN cognition and the right-side stack.
    hub_x = 11.25
    hub_y = 4.9
    ax.add_patch(FancyBboxPatch(
        (hub_x - 0.36, hub_y - 0.22), 0.72, 0.44,
        boxstyle="round,pad=0.0,rounding_size=0.08",
        facecolor="#1F2937", edgecolor="#1F2937", zorder=5,
    ))
    ax.text(hub_x, hub_y, "result\nevent",
            color="#FFFFFF", fontsize=7.5, weight="bold",
            ha="center", va="center", family="DejaVu Sans", zorder=6,
            linespacing=1.05)
    # cognition → hub
    _arrow(ax, (11.0, 4.9), (hub_x - 0.36, hub_y),
           color="#374151", lw=2.2)
    # hub → synthesis (up-right)
    _arrow(ax, (hub_x, hub_y + 0.22), (11.7, 5.6),
           color=L3["border"], lw=1.8)
    # hub → interface (down-right)
    _arrow(ax, (hub_x, hub_y - 0.22), (11.7, 4.0),
           color=L4["border"], lw=1.8)
    # synthesis → interface (summary/advice flows down into the UI)
    _arrow(ax, (13.4, 5.2), (13.4, 4.8),
           color="#374151", lw=1.6, style="-|>")

    # Bottom legend strip
    legend_y = 0.15
    ax.text(0.4, legend_y, "Built with:",
            fontsize=9, color="#8A94A6", weight="bold",
            family="DejaVu Sans")
    stack = "Streamlit · PyTorch · HuggingFace Transformers · librosa · OpenCV · Plotly · matplotlib"
    ax.text(2.05, legend_y, stack,
            fontsize=9, color="#4A5568", family="DejaVu Sans")

    OUT_PNG.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(OUT_PNG, bbox_inches="tight", facecolor="#FFFFFF",
                pad_inches=0.2)
    plt.close(fig)

    # Mirror to docs/ for the report
    OUT_PNG2.parent.mkdir(parents=True, exist_ok=True)
    OUT_PNG2.write_bytes(OUT_PNG.read_bytes())
    print(f"Saved: {OUT_PNG}")
    print(f"Saved: {OUT_PNG2}")


# ======================================================================
# 2)  python-pptx editable PPTX renderer
# ======================================================================

def _hex(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _add_rect(slide, x, y, w, h, *, fill, line, line_width=1.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.10   # corner radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex(fill)
    shape.line.color.rgb = _hex(line)
    shape.line.width = Pt(line_width)
    shape.shadow.inherit = False
    return shape


def _set_text(shape, text, *, size=12, bold=False, color="#1A1F2C",
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(40_000)
    tf.margin_right = Emu(40_000)
    tf.margin_top = Emu(20_000)
    tf.margin_bottom = Emu(20_000)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = ""
    run = p.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _hex(color)


def _add_layer_card_pptx(slide, x, y, w, h, *, spec):
    """Outer container with a small pill label inside the top-left."""
    outer = _add_rect(slide, x, y, w, h,
                      fill=spec["fill"], line=spec["border"],
                      line_width=2.0)
    # Title pill (small filled rounded rect)
    pill_w = Inches(2.3)
    pill_h = Inches(0.32)
    pill = _add_rect(slide,
                     x + Inches(0.15), y + Inches(0.15),
                     pill_w, pill_h,
                     fill=spec["border"], line=spec["border"])
    _set_text(pill, spec["label"],
              size=10, bold=True, color="#FFFFFF",
              align=PP_ALIGN.CENTER)
    # Section title under the pill
    title_box = slide.shapes.add_textbox(
        x + Inches(0.15), y + Inches(0.55),
        w - Inches(0.30), Inches(0.40),
    )
    _set_text(title_box, spec["title"],
              size=13, bold=True, color="#1A1F2C",
              anchor=MSO_ANCHOR.TOP)
    return outer


def _add_processing_card_pptx(slide, x, y, w, h, *, items):
    """Inner white card with a vertical list of items."""
    _add_rect(slide, x, y, w, h, fill="#FFFFFF", line="#C8CFDA",
              line_width=0.75)
    # PROCESSING tag
    tag = slide.shapes.add_textbox(
        x + Inches(0.15), y + Inches(0.10),
        w - Inches(0.30), Inches(0.30),
    )
    _set_text(tag, "⚙  PROCESSING",
              size=9, bold=True, color="#8A94A6",
              anchor=MSO_ANCHOR.TOP)
    # Rows
    row_h = Inches(0.42)
    base_y = y + Inches(0.50)
    for i, item in enumerate(items):
        row_y = base_y + i * row_h
        if isinstance(item, tuple):
            icon, label = item
            text = f"{icon}   {label}"
        else:
            text = item
        row = slide.shapes.add_textbox(
            x + Inches(0.20), row_y,
            w - Inches(0.40), row_h,
        )
        _set_text(row, text,
                  size=11, color="#1A1F2C",
                  anchor=MSO_ANCHOR.MIDDLE)


def _add_input_node_pptx(slide, x, y, w, h, icon, label, *, accent="#3B82F6"):
    _add_rect(slide, x, y, w, h, fill="#FFFFFF", line=accent, line_width=1.0)
    box = slide.shapes.add_textbox(x, y, w, h)
    _set_text(box, f"{icon}   {label}",
              size=11, color="#1A1F2C",
              anchor=MSO_ANCHOR.MIDDLE)


def _add_arrow_pptx(slide, x1, y1, x2, y2, *, color="#374151", weight=1.5):
    line = slide.shapes.add_connector(1, x1, y1, x2, y2)   # 1 = straight
    line.line.color.rgb = _hex(color)
    line.line.width = Pt(weight)
    # Add arrowhead (tail-less, head only)
    lnElem = line.line._get_or_add_ln()
    from lxml import etree
    nsmap = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    headEnd = etree.SubElement(
        lnElem, "{http://schemas.openxmlformats.org/drawingml/2006/main}headEnd"
    )
    headEnd.set("type", "none")
    tailEnd = etree.SubElement(
        lnElem, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd"
    )
    tailEnd.set("type", "triangle")
    tailEnd.set("w", "med")
    tailEnd.set("len", "med")


def render_pptx():
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    # White background
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0,
        prs.slide_width, prs.slide_height,
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = _hex("#FFFFFF")
    bg.line.fill.background()
    bg.shadow.inherit = False

    # Title
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.6),
    )
    _set_text(title_box, TITLE,
              size=22, bold=True, color="#1A1F2C",
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    subtitle_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.85), Inches(12.3), Inches(0.35),
    )
    _set_text(
        subtitle_box,
        "Image + Text + Audio + Video → fusion + mismatch → summary + coach → UI",
        size=12, color="#4A5568",
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )

    # --- Input nodes (left edge) ---
    for i, (icon, label) in enumerate(PERCEPTION_NODES):
        _add_input_node_pptx(
            slide,
            Inches(0.35), Inches(1.55 + i * 0.55),
            Inches(1.85), Inches(0.45),
            icon, label,
            accent=LAYERS["perception"]["border"],
        )

    # --- Layer 1 PERCEPTION ---
    L1 = LAYERS["perception"]
    L1_x, L1_y, L1_w, L1_h = Inches(2.45), Inches(1.4), Inches(3.1), Inches(5.5)
    _add_layer_card_pptx(slide, L1_x, L1_y, L1_w, L1_h, spec=L1)
    _add_processing_card_pptx(
        slide,
        L1_x + Inches(0.15), L1_y + Inches(1.05),
        L1_w - Inches(0.30), L1_h - Inches(1.20),
        items=PERCEPTION_PIPELINE,
    )

    # --- Layer 2 COGNITION ---
    L2 = LAYERS["cognition"]
    L2_x, L2_y, L2_w, L2_h = Inches(5.85), Inches(1.4), Inches(3.4), Inches(5.5)
    _add_layer_card_pptx(slide, L2_x, L2_y, L2_w, L2_h, spec=L2)
    _add_processing_card_pptx(
        slide,
        L2_x + Inches(0.15), L2_y + Inches(1.05),
        L2_w - Inches(0.30), L2_h - Inches(1.20),
        items=COGNITION_PIPELINE,
    )

    # --- Layer 3 SYNTHESIS (top-right) ---
    L3 = LAYERS["synthesis"]
    L3_x, L3_y, L3_w, L3_h = Inches(9.55), Inches(1.4), Inches(3.4), Inches(2.7)
    _add_layer_card_pptx(slide, L3_x, L3_y, L3_w, L3_h, spec=L3)
    _add_processing_card_pptx(
        slide,
        L3_x + Inches(0.15), L3_y + Inches(1.05),
        L3_w - Inches(0.30), L3_h - Inches(1.20),
        items=SYNTHESIS_PIPELINE,
    )

    # --- Layer 4 INTERFACE (bottom-right) ---
    L4 = LAYERS["interface"]
    L4_x, L4_y, L4_w, L4_h = Inches(9.55), Inches(4.2), Inches(3.4), Inches(2.7)
    _add_layer_card_pptx(slide, L4_x, L4_y, L4_w, L4_h, spec=L4)
    _add_processing_card_pptx(
        slide,
        L4_x + Inches(0.15), L4_y + Inches(1.05),
        L4_w - Inches(0.30), L4_h - Inches(1.20),
        items=INTERFACE_PIPELINE,
    )

    # --- Arrows ---
    # Inputs → Layer 1
    for i in range(len(PERCEPTION_NODES)):
        y = Inches(1.55 + i * 0.55 + 0.22)
        _add_arrow_pptx(slide, Inches(2.20), y,
                        L1_x + Inches(0.05), y,
                        color=L1["border"], weight=1.0)
    # Layer 1 → Layer 2
    _add_arrow_pptx(slide,
                    L1_x + L1_w, Inches(4.1),
                    L2_x, Inches(4.1),
                    color="#374151", weight=2.0)
    # Layer 2 → Layer 3 (event-stream split)
    _add_arrow_pptx(slide,
                    L2_x + L2_w, Inches(2.8),
                    L3_x, Inches(2.8),
                    color=L3["border"], weight=1.8)
    # Layer 2 → Layer 4
    _add_arrow_pptx(slide,
                    L2_x + L2_w, Inches(5.55),
                    L4_x, Inches(5.55),
                    color=L4["border"], weight=1.8)
    # Layer 3 → Layer 4 (summary feeds the UI)
    _add_arrow_pptx(slide,
                    L3_x + L3_w / 2, L3_y + L3_h,
                    L4_x + L4_w / 2, L4_y,
                    color="#374151", weight=1.5)

    # Footer note
    footer = slide.shapes.add_textbox(
        Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.30),
    )
    _set_text(
        footer,
        "Built with Streamlit · PyTorch · HuggingFace Transformers · librosa · OpenCV · Plotly  ·  Editable in PowerPoint",
        size=10, color="#8A94A6",
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )

    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    print(f"Saved: {OUT_PPTX}")


# ======================================================================
# Main
# ======================================================================

if __name__ == "__main__":
    render_png()
    render_pptx()
