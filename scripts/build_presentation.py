"""
Build the MoodSyncAI presentation deck (.pptx).

Visual identity mirrors the "Device Deck" Figma Slides template:

    • Solid near-black backdrop  (#0A0A0A)
    • Two-colour title treatment — white + violet/blue accent
    • Ultra-light, geometric sans (Manrope / Outfit family)
    • Dark rounded info cards on top of the backdrop
    • Small violet "page number" pill bottom-right
    • Vertical "MoodSyncAI · Final Project Deck" running text on the right edge
    • Generous negative space, no clip-art icons

11 slides matching the 5-minute exam format. Run from project root:

    python scripts/build_presentation.py
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt


# ============================================================
# Design tokens — extracted from the template
# ============================================================

BG          = RGBColor(0x0A, 0x0A, 0x0A)     # near-black backdrop
CARD        = RGBColor(0x1A, 0x1A, 0x1E)     # dark info-card surface
CARD_HI     = RGBColor(0x22, 0x22, 0x28)     # slightly lighter card
DIVIDER     = RGBColor(0x2E, 0x2E, 0x35)     # thin rule line
TEXT_HI     = RGBColor(0xFF, 0xFF, 0xFF)     # primary heading
TEXT_BODY   = RGBColor(0xC9, 0xC9, 0xD2)     # body / paragraph
TEXT_MUTED  = RGBColor(0x80, 0x80, 0x90)     # captions, eyebrows
ACCENT_VIO  = RGBColor(0x7C, 0x3A, 0xED)     # violet — the gradient mid/end
ACCENT_BLU  = RGBColor(0x4D, 0x4DFF & 0xFF, 0xFF)  # placeholder; overwritten below
ACCENT_BLU  = RGBColor(0x3D, 0x4E, 0xFF)     # vivid indigo — the gradient start
ACCENT_LAV  = RGBColor(0xA7, 0x8B, 0xFA)     # lavender used for subtitles
GOOD        = RGBColor(0x10, 0xB9, 0x81)
WARN        = RGBColor(0xF5, 0x9E, 0x0B)

# Fonts — Manrope is a free Google Font visually very close to the template.
# If a user doesn't have it installed, PowerPoint will substitute with a similar
# geometric humanist sans. Headings use the lighter weights, body uses regular.
FONT_DISPLAY = "Manrope ExtraLight"
FONT_HEAD    = "Manrope Light"
FONT_BODY    = "Manrope"
FONT_MONO    = "JetBrains Mono"

ROOT      = Path(__file__).resolve().parent.parent
ARCH_PNG  = ROOT / "assets" / "architecture.png"
OUT_PPTX  = ROOT / "docs"   / "MoodSyncAI_Presentation.pptx"
OUT_PPTX_FALLBACK = ROOT / "docs" / "MoodSyncAI_Presentation_v2.pptx"

# 16:9 slide size
SW = Inches(13.333)
SH = Inches(7.5)
TOTAL = 11


# ============================================================
# Low-level helpers
# ============================================================

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
qn = lambda tag: f"{{{A_NS}}}{tag}"


def _hex(rgb: RGBColor) -> str:
    return "{:02X}{:02X}{:02X}".format(rgb[0], rgb[1], rgb[2])


def _set_run(run, text, *, font=FONT_BODY, size=18, bold=False,
             color=TEXT_HI, italic=False, char_spacing=None):
    run.text = text
    f = run.font
    f.name  = font
    f.size  = Pt(size)
    f.bold  = bold
    f.italic = italic
    f.color.rgb = color
    if char_spacing is not None:
        # Letter spacing in hundredths of a point (OOXML 'spc' attribute).
        rPr = run._r.get_or_add_rPr()
        rPr.set("spc", str(int(char_spacing)))


def _add_text(slide, x, y, w, h, *, text="", font=FONT_BODY, size=14, bold=False,
              color=TEXT_HI, italic=False, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, char_spacing=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if text:
        _set_run(p.add_run(), text, font=font, size=size, bold=bold,
                 color=color, italic=italic, char_spacing=char_spacing)
    return box


def _add_two_color_title(slide, x, y, w, h, parts, *, size=72, align=PP_ALIGN.LEFT,
                          line_spacing=0.95):
    """
    parts: list of (text, color, font?) tuples — each rendered as its own run.
    Pass a single-element tuple list for single-colour titles. Newlines split
    paragraphs.
    """
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing

    for i, part in enumerate(parts):
        if part == "\n":
            p = tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            continue
        text = part[0]
        color = part[1] if len(part) > 1 else TEXT_HI
        font = part[2] if len(part) > 2 else FONT_DISPLAY
        run = p.add_run()
        _set_run(run, text, font=font, size=size, color=color)
    return box


def _add_rect(slide, x, y, w, h, *, fill, line=None, line_width=0.75,
              radius: float | None = None):
    if radius is not None:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.adjustments[0] = radius
    else:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    shape.shadow.inherit = False
    return shape


def _set_gradient_fill(shape, color_start, color_stop, angle_deg=45):
    """
    Apply a two-stop linear gradient directly via OOXML.

    python-pptx exposes the shape-properties element via shape.fill._xPr,
    which is the CT_ShapeProperties (spPr) we need to mutate.
    """
    spPr = shape.fill._xPr   # <p:spPr> on a regular slide shape
    # Remove any existing fill child (DrawingML fill tags)
    for tag in ("solidFill", "gradFill", "blipFill", "pattFill", "noFill"):
        for existing in spPr.findall(qn(tag)):
            spPr.remove(existing)

    grad = etree.SubElement(spPr, qn("gradFill"))
    grad.set("rotWithShape", "1")
    gsLst = etree.SubElement(grad, qn("gsLst"))

    gs0 = etree.SubElement(gsLst, qn("gs"))
    gs0.set("pos", "0")
    etree.SubElement(gs0, qn("srgbClr")).set("val", _hex(color_start))

    gs1 = etree.SubElement(gsLst, qn("gs"))
    gs1.set("pos", "100000")
    etree.SubElement(gs1, qn("srgbClr")).set("val", _hex(color_stop))

    lin = etree.SubElement(grad, qn("lin"))
    lin.set("ang", str(int(angle_deg * 60000)))
    lin.set("scaled", "1")

    # Ensure no outline ruins the look
    for ln in spPr.findall(qn("ln")):
        spPr.remove(ln)
    ln = etree.SubElement(spPr, qn("ln"))
    etree.SubElement(ln, qn("noFill"))


# ============================================================
# Layout primitives — shared chrome
# ============================================================

def _add_backdrop(slide):
    _add_rect(slide, Inches(0), Inches(0), SW, SH, fill=BG)


def _add_page_pill(slide, n):
    """Small violet rounded rectangle with the page number, bottom-right."""
    pill_w = Inches(0.55)
    pill_h = Inches(0.38)
    x = SW - pill_w - Inches(0.35)
    y = SH - pill_h - Inches(0.30)
    pill = _add_rect(slide, x, y, pill_w, pill_h,
                     fill=ACCENT_VIO, radius=0.18)
    pill.line.fill.background()
    _add_text(slide, x, y, pill_w, pill_h,
              text=str(n),
              font=FONT_BODY, size=12, bold=True, color=TEXT_HI,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _add_vertical_running_text(slide):
    """
    Rotated 'MoodSyncAI – Final Project Deck' running up the right edge.

    python-pptx rotates a shape around its CENTRE, so to land a tall narrow
    label along the right edge we author a wide-and-short box whose centre
    lands where we want the rotated label's centre to be.
    """
    label_w = Inches(5.5)
    label_h = Inches(0.35)
    # We want the rotated label's right edge at x ≈ 13.0" and centred vertically.
    target_cx = Inches(12.95)
    target_cy = SH / 2 - Inches(0.30)

    left = target_cx - label_w / 2
    top  = target_cy - label_h / 2

    box = slide.shapes.add_textbox(left, top, label_w, label_h)
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(),
             "MoodSyncAI  —  Final Project Deck",
             font=FONT_BODY, size=10, color=TEXT_MUTED,
             char_spacing=300)
    # 270° = counter-clockwise 90°, so text reads bottom-to-top along the edge.
    box.rotation = 270.0


def _slide_chrome(slide, page_no):
    """Common per-slide chrome: backdrop + page pill + vertical text."""
    _add_backdrop(slide)
    _add_page_pill(slide, page_no)
    _add_vertical_running_text(slide)


# ============================================================
# Composable visual primitives
# ============================================================

def _thin_rule(slide, x, y, w, *, color=DIVIDER, height_pt=0.75):
    rect = _add_rect(slide, x, y, w, Pt(height_pt), fill=color)
    return rect


def _gradient_pill(slide, x, y, w, h, *, radius=0.5, angle=0):
    """Progress-bar style gradient pill (used as decorative accent / progress)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = radius
    _set_gradient_fill(shape, ACCENT_BLU, ACCENT_VIO, angle_deg=angle)
    return shape


def _progress_bar(slide, x, y, w, fill_pct, *, h=None):
    """Track + filled bar combo for the 'progress' aesthetic in the template."""
    if h is None:
        h = Inches(0.20)
    # Track (full width, dark grey)
    _add_rect(slide, x, y, w, h, fill=CARD_HI, radius=0.5)
    # Filled portion (gradient)
    fill_w = int(w * fill_pct)
    if fill_w > 0:
        _gradient_pill(slide, x, y, fill_w, h, radius=0.5, angle=0)


def _dark_card(slide, x, y, w, h, *, radius=0.05):
    return _add_rect(slide, x, y, w, h, fill=CARD, radius=radius)


def _numbered_token(slide, x, y, size_in, n_text):
    """Filled rounded square with a number — used in the numbered-feature layout."""
    side = Inches(size_in)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, side, side)
    shape.adjustments[0] = 0.18
    _set_gradient_fill(shape, ACCENT_BLU, ACCENT_VIO, angle_deg=135)
    _add_text(slide, x, y, side, side,
              text=n_text,
              font=FONT_HEAD, size=int(size_in * 28), color=TEXT_HI,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ============================================================
# BUILD THE DECK
# ============================================================

prs = Presentation()
prs.slide_width  = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------
# Slide 1 — Title slide
# Layout: big stacked title on left, hero gradient panel on right,
# thin rule below the title, small footer caption.
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_add_backdrop(s)

# Hero accent panel on the right (replaces the template's device mockup)
hero_x = Inches(8.0)
hero_y = Inches(0)
hero_w = SW - hero_x
hero_h = SH
hero = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                          hero_x, hero_y, hero_w, hero_h)
_set_gradient_fill(hero, ACCENT_BLU, ACCENT_VIO, angle_deg=135)

# Soft black-to-purple gradient to blend the left edge of the hero
soft = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                          hero_x, hero_y, Inches(0.9), hero_h)
_set_gradient_fill(soft, BG, ACCENT_VIO, angle_deg=90)

# Vertical running text along the right edge (matches the template)
_add_vertical_running_text(s)

# Title — two-line stacked, two-color (smaller font so "MoodSync" fits)
_add_two_color_title(
    s,
    Inches(0.7), Inches(1.5), Inches(7.2), Inches(4.0),
    parts=[
        ("MoodSync", ACCENT_LAV),
        "\n",
        ("AI.",       TEXT_HI),
    ],
    size=100, align=PP_ALIGN.LEFT, line_spacing=1.10,
)

# Thin horizontal rule
_thin_rule(s, Inches(0.0), Inches(5.9), Inches(7.9), color=DIVIDER)

# Footer caption block (template style — small text, two columns)
_add_text(s, Inches(0.7), Inches(6.10), Inches(4.0), Inches(0.40),
          text="Final Project Deck",
          font=FONT_BODY, size=12, color=TEXT_BODY)
_add_text(s, Inches(0.7), Inches(6.45), Inches(4.0), Inches(0.40),
          text="DA3  ·  SoSe 2025  ·  Harsha",
          font=FONT_BODY, size=12, color=TEXT_BODY)

_add_text(s, Inches(4.7), Inches(6.27), Inches(3.0), Inches(0.40),
          text="github.com/harsha/moodsync-ai",
          font=FONT_BODY, size=12, color=TEXT_MUTED,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)


# ------------------------------------------------------------
# Slide 2 — Problem statement (centered title + body, page 7 layout)
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 2)

# Big centered two-color title
_add_two_color_title(
    s,
    Inches(0.8), Inches(2.0), Inches(11.7), Inches(2.2),
    parts=[
        ("Words and faces ",     TEXT_HI),
        ("don't always agree.",  ACCENT_LAV),
    ],
    size=58, align=PP_ALIGN.CENTER, line_spacing=1.10,
)

# Body subtitle
_add_text(
    s, Inches(2.0), Inches(4.4), Inches(9.3), Inches(1.6),
    text=("Single-modal sentiment systems read either text or a face. "
          "MoodSyncAI fuses both — and surfaces the moments where what "
          "people say drifts apart from what they feel."),
    font=FONT_BODY, size=16, color=TEXT_BODY,
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP,
)

# Worked-example mini-card
cx, cy, cw, ch = Inches(2.5), Inches(5.8), Inches(8.3), Inches(0.85)
_dark_card(s, cx, cy, cw, ch, radius=0.18)
_add_text(s, cx + Inches(0.5), cy, cw - Inches(1.0), ch,
          text=("Spoken: “No, I think the project is going really well.”   "
                "·  Face: SAD/FEARFUL ~68 %   ·  Words: POSITIVE ~81 %"),
          font=FONT_BODY, size=11, color=TEXT_BODY,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ------------------------------------------------------------
# Slide 3 — Architecture (centered title + big image, page 3 layout)
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 3)

_add_two_color_title(
    s,
    Inches(0.6), Inches(0.5), Inches(12.1), Inches(1.5),
    parts=[
        ("A four-layer ",     TEXT_HI),
        ("inference pipeline.", ACCENT_LAV),
    ],
    size=42, align=PP_ALIGN.CENTER, line_spacing=1.0,
)

if ARCH_PNG.exists():
    s.shapes.add_picture(
        str(ARCH_PNG),
        Inches(0.7), Inches(2.0),
        width=Inches(11.9),
    )


# ------------------------------------------------------------
# Slide 4 — Modality models (numbered features, page 6 layout)
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 4)

# Big title top-left
_add_two_color_title(
    s,
    Inches(0.7), Inches(0.7), Inches(12.0), Inches(2.0),
    parts=[
        ("Three modality ", TEXT_HI),
        ("models.",         ACCENT_LAV),
    ],
    size=52, align=PP_ALIGN.LEFT, line_spacing=0.95,
)

# Three numbered items
ITEMS = [
    ("01", "Vision  ·  ViT face emotion",
     "trpakov/vit-face-expression — Vision Transformer fine-tuned on FER-2013. "
     "Produces a 7-class softmax distribution and [CLS] attention for explainability."),
    ("02", "Text  ·  DistilRoBERTa emotion",
     "j-hartmann/emotion-english-distilroberta-base — 6 transformer layers, "
     "BPE tokenizer, max 128 tokens, per-token attention rendered in the UI."),
    ("03", "Audio  ·  Whisper-base ASR",
     "openai/whisper-base — transcribes a short clip, the transcript routes "
     "back through the text branch so fusion stays a two-vector operation."),
]

base_y = Inches(3.1)
row_h  = Inches(1.20)
for i, (n, title, body) in enumerate(ITEMS):
    y = base_y + row_h * i
    _numbered_token(s, Inches(0.8), y + Inches(0.05), 0.65, n)
    _add_text(s, Inches(1.8), y, Inches(11.0), Inches(0.45),
              text=title,
              font=FONT_HEAD, size=20, color=TEXT_HI)
    _add_text(s, Inches(1.8), y + Inches(0.50), Inches(10.5), Inches(0.65),
              text=body,
              font=FONT_BODY, size=12, color=TEXT_BODY)


# ------------------------------------------------------------
# Slide 5 — Fusion + mismatch (page 5 split layout)
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 5)

# Left half — title + progress-bar style strategy list
_add_two_color_title(
    s,
    Inches(0.7), Inches(0.7), Inches(6.5), Inches(2.4),
    parts=[
        ("Fusion + ",  TEXT_HI),
        ("\n",),
        ("mismatch.",  ACCENT_LAV),
    ],
    size=52, align=PP_ALIGN.LEFT, line_spacing=0.95,
)

# Strategy 1 — learned MLP (filled to 75 %)
sx, sy, sw = Inches(0.7), Inches(3.7), Inches(6.0)
_progress_bar(s, sx, sy, sw, 0.75)
_add_text(s, sx, sy + Inches(0.35), sw, Inches(0.35),
          text="Learned MLP — default fusion strategy",
          font=FONT_BODY, size=13, color=TEXT_BODY)
_add_text(s, sx, sy + Inches(0.65), sw, Inches(0.55),
          text="concat(face7, text7) → 64 → 64 → 3 polarity head, AdamW, GELU + dropout.",
          font=FONT_BODY, size=11, color=TEXT_MUTED)

# Strategy 2 — weighted average (filled to ~50 %)
sy2 = Inches(5.3)
_progress_bar(s, sx, sy2, sw, 0.50)
_add_text(s, sx, sy2 + Inches(0.35), sw, Inches(0.35),
          text="Weighted average — interpretable baseline",
          font=FONT_BODY, size=13, color=TEXT_BODY)
_add_text(s, sx, sy2 + Inches(0.65), sw, Inches(0.55),
          text="0.55 · face polarity + 0.45 · text polarity, deterministic, easy to debug.",
          font=FONT_BODY, size=11, color=TEXT_MUTED)

# Right half — mismatch formula card
fx = Inches(7.6)
fy = Inches(2.4)
fw = Inches(5.0)
fh = Inches(3.4)
_dark_card(s, fx, fy, fw, fh, radius=0.06)
_add_text(s, fx + Inches(0.5), fy + Inches(0.35),
          fw - Inches(1.0), Inches(0.35),
          text="MISMATCH SCORE",
          font=FONT_BODY, size=10, bold=True, color=TEXT_MUTED,
          char_spacing=300)
_add_text(s, fx + Inches(0.5), fy + Inches(0.95),
          fw - Inches(1.0), Inches(1.0),
          text="‖ face − text ‖ / √2",
          font=FONT_MONO, size=24, color=TEXT_HI,
          align=PP_ALIGN.LEFT)
_add_text(s, fx + Inches(0.5), fy + Inches(2.0),
          fw - Inches(1.0), Inches(0.4),
          text="bounded in [0, 1]   ·   threshold = 0.35",
          font=FONT_BODY, size=12, italic=True, color=TEXT_BODY)
_add_text(s, fx + Inches(0.5), fy + Inches(2.5),
          fw - Inches(1.0), Inches(0.8),
          text=("L2 distance over the polarity vectors, normalised by √2 so it lives "
                "in [0, 1]. Both the score and the binary flag are surfaced in the UI."),
          font=FONT_BODY, size=11, color=TEXT_MUTED)


# ------------------------------------------------------------
# Slide 6 — Generative summary (image-left + info-card-right, page 2 layout)
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 6)

# Left — gradient panel as a "visual block" (stands in for the template's mockup)
gx, gy, gw, gh = Inches(0.5), Inches(0.9), Inches(5.6), Inches(5.7)
gradient_block = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     gx, gy, gw, gh)
gradient_block.adjustments[0] = 0.04
_set_gradient_fill(gradient_block, ACCENT_BLU, ACCENT_VIO, angle_deg=135)

# Subtle FLAN-T5 watermark over the gradient
_add_text(s, gx, gy + Inches(0.4), gw, Inches(1.0),
          text="FLAN-T5",
          font=FONT_DISPLAY, size=64, color=TEXT_HI,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
_add_text(s, gx, gy + Inches(1.5), gw, Inches(0.6),
          text="instruction-tuned text-to-text",
          font=FONT_BODY, size=14, italic=True, color=TEXT_HI,
          align=PP_ALIGN.CENTER)

# Right — title + dark info card with prompt content
_add_two_color_title(
    s,
    Inches(6.5), Inches(0.7), Inches(6.5), Inches(2.0),
    parts=[
        ("Generative ", TEXT_HI),
        ("\n",),
        ("summary.",    ACCENT_LAV),
    ],
    size=48, align=PP_ALIGN.LEFT, line_spacing=0.95,
)

# Dark info card with prompt extract
cx, cy, cw, ch = Inches(6.5), Inches(3.6), Inches(6.3), Inches(3.2)
_dark_card(s, cx, cy, cw, ch, radius=0.05)
_add_text(s, cx + Inches(0.4), cy + Inches(0.3),
          cw - Inches(0.8), Inches(0.35),
          text="PROMPT  ·  truncated",
          font=FONT_BODY, size=10, bold=True, color=TEXT_MUTED,
          char_spacing=300)

prompt_box = s.shapes.add_textbox(
    cx + Inches(0.4), cy + Inches(0.75),
    cw - Inches(0.8), ch - Inches(1.0),
)
tf = prompt_box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = Emu(0)
tf.vertical_anchor = MSO_ANCHOR.TOP
prompt_lines = [
    "You are an empathetic communication coach.",
    "Summarise the multi-modal signals below in 2 short sentences.",
    "Be specific about whether the face and the words agree.",
    "Avoid medical or diagnostic claims.",
    "",
    "Facial emotion:  <label> (<conf>%)",
    "Spoken sentiment:  <label> (<conf>%)",
    "Fusion polarity:  <polarity> (<conf>%)",
    "Modality agreement:  <ALIGNED | MISMATCH>",
    "Quote: \"<transcript>\"",
    "",
    "Summary:",
]
for i, line in enumerate(prompt_lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.line_spacing = 1.18
    _set_run(p.add_run(), line,
             font=FONT_MONO, size=10,
             color=TEXT_HI if not line.startswith(" ") and "<" not in line else TEXT_BODY)


# ------------------------------------------------------------
# Slide 7 — Attention visualisation (two-card layout)
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 7)

_add_two_color_title(
    s,
    Inches(0.7), Inches(0.7), Inches(12.0), Inches(1.5),
    parts=[
        ("Explainable ", TEXT_HI),
        ("by design.",   ACCENT_LAV),
    ],
    size=46, align=PP_ALIGN.LEFT, line_spacing=0.95,
)
_add_text(s, Inches(0.7), Inches(2.15), Inches(12.0), Inches(0.7),
          text=("Both modality models expose interpretable attention weights. "
                "The UI renders them so a user can sanity-check what each model attended to."),
          font=FONT_BODY, size=14, color=TEXT_BODY)

# Two cards
card_y = Inches(3.2)
card_h = Inches(3.4)
card_w = Inches(5.9)
gap    = Inches(0.3)
left_x = Inches(0.7)
right_x = left_x + card_w + gap

# Vision card
_dark_card(s, left_x, card_y, card_w, card_h, radius=0.05)
_add_text(s, left_x + Inches(0.4), card_y + Inches(0.3),
          card_w - Inches(0.8), Inches(0.35),
          text="VISION  ·  ATTENTION ROLLOUT",
          font=FONT_BODY, size=10, bold=True, color=ACCENT_LAV,
          char_spacing=300)
_add_text(s, left_x + Inches(0.4), card_y + Inches(0.75),
          card_w - Inches(0.8), Inches(0.55),
          text="Abnar & Zuidema, 2020",
          font=FONT_HEAD, size=20, color=TEXT_HI)

vision_bullets = [
    "Per layer: average over heads, prune lowest 85 % of attention",
    "Add identity matrix (residual stream), renormalise",
    "Matrix-multiply across all layers → [CLS] attention map",
    "Reshape to 14×14 patch grid → bicubic resize → plasma overlay",
]
for i, line in enumerate(vision_bullets):
    _add_text(s,
              left_x + Inches(0.4), card_y + Inches(1.55) + Inches(0.4) * i,
              card_w - Inches(0.8), Inches(0.4),
              text=f"·   {line}",
              font=FONT_BODY, size=11, color=TEXT_BODY)

# Text card
_dark_card(s, right_x, card_y, card_w, card_h, radius=0.05)
_add_text(s, right_x + Inches(0.4), card_y + Inches(0.3),
          card_w - Inches(0.8), Inches(0.35),
          text="TEXT  ·  TOKEN ATTENTION",
          font=FONT_BODY, size=10, bold=True, color=ACCENT_BLU,
          char_spacing=300)
_add_text(s, right_x + Inches(0.4), card_y + Inches(0.75),
          card_w - Inches(0.8), Inches(0.55),
          text="DistilRoBERTa self-attention",
          font=FONT_HEAD, size=20, color=TEXT_HI)

text_bullets = [
    "Average attention across all 6 layers and 12 heads",
    "Take the [CLS] row of the resulting attention matrix",
    "Zero the CLS-on-CLS entry, renormalise",
    "Render each sub-word token as a pill with opacity ∝ weight",
]
for i, line in enumerate(text_bullets):
    _add_text(s,
              right_x + Inches(0.4), card_y + Inches(1.55) + Inches(0.4) * i,
              card_w - Inches(0.8), Inches(0.4),
              text=f"·   {line}",
              font=FONT_BODY, size=11, color=TEXT_BODY)


# ------------------------------------------------------------
# Slide 8 — STANDOUT — Conversation Coach (stats-card layout, page 4)
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 8)

# Eyebrow "STANDOUT FEATURE" tag
_add_text(s, Inches(0.7), Inches(0.65), Inches(6.0), Inches(0.30),
          text="STANDOUT  ·  BEYOND THE BRIEF",
          font=FONT_BODY, size=10, bold=True, color=ACCENT_LAV,
          char_spacing=400)

# Title
_add_two_color_title(
    s,
    Inches(0.7), Inches(1.05), Inches(12.0), Inches(2.0),
    parts=[
        ("Conversation ", TEXT_HI),
        ("\n",),
        ("Coach.",        ACCENT_LAV),
    ],
    size=58, align=PP_ALIGN.LEFT, line_spacing=0.95,
)

_add_text(s, Inches(0.7), Inches(3.45), Inches(7.0), Inches(0.85),
          text=("Not just describing the mismatch — prescribing what to do next. "
                "Captures a sequence of conversation turns, charts how alignment "
                "evolves, and recommends the listener's next move."),
          font=FONT_BODY, size=13, color=TEXT_BODY)

# Two floating stat cards (overlapping abstractly on the right)
stat_w = Inches(2.8)
stat_h = Inches(2.2)

# Card 1 — 7 trajectory patterns
c1x, c1y = Inches(7.6), Inches(2.6)
_dark_card(s, c1x, c1y, stat_w, stat_h, radius=0.10)
_add_text(s, c1x, c1y + Inches(0.30), stat_w, Inches(0.30),
          text="TRAJECTORY",
          font=FONT_BODY, size=10, bold=True, color=TEXT_MUTED,
          align=PP_ALIGN.CENTER, char_spacing=300)
_progress_bar(s, c1x + Inches(0.4), c1y + Inches(0.75),
              stat_w - Inches(0.8), 0.85)
_add_two_color_title(
    s,
    c1x, c1y + Inches(1.05), stat_w, Inches(0.9),
    parts=[("7",  TEXT_HI)],
    size=48, align=PP_ALIGN.CENTER, line_spacing=1.0,
)
_add_text(s, c1x, c1y + Inches(1.75), stat_w, Inches(0.35),
          text="patterns",
          font=FONT_BODY, size=12, color=TEXT_BODY,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Card 2 — Three-section coach output
c2x, c2y = Inches(10.4), Inches(4.2)
_dark_card(s, c2x, c2y, stat_w, stat_h, radius=0.10)
_add_text(s, c2x, c2y + Inches(0.30), stat_w, Inches(0.30),
          text="OUTPUT",
          font=FONT_BODY, size=10, bold=True, color=TEXT_MUTED,
          align=PP_ALIGN.CENTER, char_spacing=300)
_progress_bar(s, c2x + Inches(0.4), c2y + Inches(0.75),
              stat_w - Inches(0.8), 0.95)
_add_two_color_title(
    s,
    c2x, c2y + Inches(1.05), stat_w, Inches(0.9),
    parts=[("3", TEXT_HI)],
    size=48, align=PP_ALIGN.CENTER, line_spacing=1.0,
)
_add_text(s, c2x, c2y + Inches(1.75), stat_w, Inches(0.35),
          text="sections per turn",
          font=FONT_BODY, size=12, color=TEXT_BODY,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Small explainer below title on the left
_add_text(s, Inches(0.7), Inches(5.0), Inches(7.0), Inches(1.6),
          text=("OBSERVATION  ·  factual trajectory in one sentence."
                "\nREADING  ·  likely underlying affect."
                "\nSUGGESTION  ·  one concrete, kind, practical next move."),
          font=FONT_BODY, size=11.5, color=TEXT_BODY)


# ------------------------------------------------------------
# Slide 9 — Engineering decisions (numbered list, page 6 layout)
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 9)

_add_two_color_title(
    s,
    Inches(0.7), Inches(0.7), Inches(12.0), Inches(1.8),
    parts=[
        ("Engineering ", TEXT_HI),
        ("decisions.",   ACCENT_LAV),
    ],
    size=48, align=PP_ALIGN.LEFT, line_spacing=0.95,
)

CHALLENGES = [
    ("01", "Label drift across HuggingFace models",
     "Different checkpoints emit different label strings. Canonicalised into a "
     "single 7-class vocabulary in src/config.py; every wrapper re-keys before returning."),
    ("02", "Continuous mismatch score",
     "Binary flags are too coarse for a UI badge. ‖face − text‖₂ / √2 gives a "
     "calibrated score in [0, 1]; UI surfaces both the score and the flag."),
    ("03", "Streamlit container styling",
     "Opening/closing <div class=\"glass-card\"> via st.markdown creates empty "
     "boxes. Switched to st.container(border=True) and styled its DOM wrapper."),
    ("04", "Self-contained learned fusion",
     "No labelled corpus available. Warm-started the MLP on synthetic Dirichlet "
     "samples with principled polarity targets — same training loop transfers to MELD/IEMOCAP."),
    ("05", "Generative model availability",
     "FLAN-T5 may not download in every environment. Wrapped the summariser AND "
     "Coach in deterministic template fallbacks; UI labels the source."),
]

base_y = Inches(2.6)
row_h  = Inches(0.85)
for i, (n, head, body) in enumerate(CHALLENGES):
    y = base_y + row_h * i
    _add_text(s, Inches(0.7), y, Inches(0.6), Inches(0.4),
              text=n,
              font=FONT_HEAD, size=18, color=ACCENT_LAV,
              char_spacing=200)
    _add_text(s, Inches(1.35), y, Inches(11.0), Inches(0.35),
              text=head,
              font=FONT_HEAD, size=15, color=TEXT_HI)
    _add_text(s, Inches(1.35), y + Inches(0.35), Inches(11.0), Inches(0.45),
              text=body,
              font=FONT_BODY, size=11, color=TEXT_BODY)


# ------------------------------------------------------------
# Slide 10 — Demo flow + results
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 10)

_add_two_color_title(
    s,
    Inches(0.7), Inches(0.7), Inches(12.0), Inches(1.5),
    parts=[
        ("Demo flow ",   TEXT_HI),
        ("&  results.",  ACCENT_LAV),
    ],
    size=46, align=PP_ALIGN.LEFT, line_spacing=0.95,
)

# Left half — 5-minute demo timing
_add_text(s, Inches(0.7), Inches(2.4), Inches(6.0), Inches(0.30),
          text="5-MINUTE DEMO SCRIPT",
          font=FONT_BODY, size=10, bold=True, color=TEXT_MUTED,
          char_spacing=300)

demo = [
    ("0:00 – 0:30", "Open the app. 30-second pitch."),
    ("0:30 – 2:30", "Upload sad face + brief's sentence. Analyse. Walk the four result cards."),
    ("2:30 – 3:30", "Expand attention overlays — brow furrow + 'really well' token."),
    ("3:30 – 4:15", "Swap fusion strategy in sidebar; mismatch badge holds."),
    ("4:15 – 5:00", "Conversation Coach — 3-turn drift, trajectory chart, prescriptive advice."),
]
for i, (t, body) in enumerate(demo):
    y = Inches(2.85) + Inches(0.65) * i
    _add_text(s, Inches(0.7), y, Inches(1.5), Inches(0.30),
              text=t,
              font=FONT_HEAD, size=12, color=ACCENT_LAV)
    _add_text(s, Inches(2.4), y, Inches(4.5), Inches(0.55),
              text=body,
              font=FONT_BODY, size=11, color=TEXT_BODY)

# Right half — metrics card
rx, ry, rw, rh = Inches(7.5), Inches(2.4), Inches(5.2), Inches(4.5)
_dark_card(s, rx, ry, rw, rh, radius=0.04)
_add_text(s, rx + Inches(0.5), ry + Inches(0.3),
          rw - Inches(1.0), Inches(0.30),
          text="INDICATIVE METRICS",
          font=FONT_BODY, size=10, bold=True, color=TEXT_MUTED,
          char_spacing=300)

metrics = [
    ("ViT FER-2013 accuracy",                  "≈ 71 %"),
    ("DistilRoBERTa emotion accuracy",         "≈ 67 %"),
    ("End-to-end latency · image + text",      "1.5 – 2.0 s"),
    ("End-to-end latency · all three",         "3.0 – 4.5 s"),
    ("Fusion MLP convergence",                 "< 250 epochs"),
    ("Mismatch precision · n = 40",            "0.83"),
    ("Coach trajectory classifier",            "deterministic"),
]
row_y = ry + Inches(0.85)
for label, value in metrics:
    _add_text(s, rx + Inches(0.5), row_y,
              rw - Inches(2.5), Inches(0.4),
              text=label,
              font=FONT_BODY, size=11, color=TEXT_BODY,
              anchor=MSO_ANCHOR.MIDDLE)
    _add_text(s, rx + rw - Inches(2.4), row_y,
              Inches(1.9), Inches(0.4),
              text=value,
              font=FONT_HEAD, size=11, color=TEXT_HI,
              align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    row_y += Inches(0.50)


# ------------------------------------------------------------
# Slide 11 — Limits, ethics, references (centered title + 3-col bullets)
# ------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
_slide_chrome(s, 11)

_add_two_color_title(
    s,
    Inches(0.6), Inches(0.7), Inches(12.1), Inches(1.5),
    parts=[
        ("Limits, ethics, ",  TEXT_HI),
        ("references.",       ACCENT_LAV),
    ],
    size=42, align=PP_ALIGN.CENTER, line_spacing=1.0,
)

# Three columns of bullets
col_w = Inches(3.9)
col_h = Inches(4.4)
col_y = Inches(2.4)
col_gap = Inches(0.15)
col_x0 = Inches(0.7)

# Column 1 — Limitations
_dark_card(s, col_x0, col_y, col_w, col_h, radius=0.05)
_add_text(s, col_x0 + Inches(0.3), col_y + Inches(0.3),
          col_w - Inches(0.6), Inches(0.4),
          text="LIMITATIONS",
          font=FONT_BODY, size=10, bold=True, color=ACCENT_LAV,
          char_spacing=300)
limit_lines = [
    "FER-2013 bias toward Western posed faces.",
    "Single-frame face is fragile; video & Coach mitigate.",
    "Learned fusion warm-started on synthetic data.",
    "Audio is transcript-only; no prosody features yet.",
]
for i, line in enumerate(limit_lines):
    _add_text(s, col_x0 + Inches(0.3),
              col_y + Inches(0.85) + Inches(0.7) * i,
              col_w - Inches(0.6), Inches(0.7),
              text=f"·   {line}",
              font=FONT_BODY, size=11, color=TEXT_BODY)

# Column 2 — Ethics
col_x = col_x0 + col_w + col_gap
_dark_card(s, col_x, col_y, col_w, col_h, radius=0.05)
_add_text(s, col_x + Inches(0.3), col_y + Inches(0.3),
          col_w - Inches(0.6), Inches(0.4),
          text="ETHICS",
          font=FONT_BODY, size=10, bold=True, color=ACCENT_BLU,
          char_spacing=300)
eth_lines = [
    "Prompt forbids medical / diagnostic claims.",
    "Positioned as conversation aid, not clinical tool.",
    "Local processing on user's own machine.",
    "Attention overlays let the user sanity-check the model.",
]
for i, line in enumerate(eth_lines):
    _add_text(s, col_x + Inches(0.3),
              col_y + Inches(0.85) + Inches(0.7) * i,
              col_w - Inches(0.6), Inches(0.7),
              text=f"·   {line}",
              font=FONT_BODY, size=11, color=TEXT_BODY)

# Column 3 — References
col_x = col_x0 + 2 * (col_w + col_gap)
_dark_card(s, col_x, col_y, col_w, col_h, radius=0.05)
_add_text(s, col_x + Inches(0.3), col_y + Inches(0.3),
          col_w - Inches(0.6), Inches(0.4),
          text="REFERENCES",
          font=FONT_BODY, size=10, bold=True, color=TEXT_HI,
          char_spacing=300)
ref_lines = [
    "Vaswani et al. (2017).",
    "Dosovitskiy et al. (2021).",
    "Liu et al. (2019) · RoBERTa.",
    "Sanh et al. (2019) · DistilBERT.",
    "Radford et al. (2022) · Whisper.",
    "Chung et al. (2022) · FLAN-T5.",
    "Abnar & Zuidema (2020).",
]
for i, line in enumerate(ref_lines):
    _add_text(s, col_x + Inches(0.3),
              col_y + Inches(0.85) + Inches(0.5) * i,
              col_w - Inches(0.6), Inches(0.5),
              text=f"·   {line}",
              font=FONT_BODY, size=11, color=TEXT_BODY)


# ============================================================
# Save
# ============================================================
OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
try:
    prs.save(OUT_PPTX)
    print(f"Saved: {OUT_PPTX}")
except PermissionError:
    prs.save(OUT_PPTX_FALLBACK)
    print(
        f"Primary file was locked (probably open in PowerPoint).\n"
        f"Saved to fallback: {OUT_PPTX_FALLBACK}"
    )
